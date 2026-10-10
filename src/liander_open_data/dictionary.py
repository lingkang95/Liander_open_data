"""Parsing of the Liander *SBI-code -> recommended profile* dictionary.

Liander publishes an Excel workbook (``Overzicht met per SBI-code een
aanbevolen profiel.xlsx``) that lists, for every SBI code (Dutch Standard
Industrial Classification), which load profile best represents businesses
in that sector.  This module turns that workbook into clean, typed
:class:`SBICode` records with English field names.

A pre-parsed copy of the dictionary ships with the package, so
:func:`load_dictionary` works without the original workbook.
"""

from __future__ import annotations

import csv
import math
from collections.abc import Iterable
from dataclasses import asdict, dataclass, fields
from importlib import resources
from pathlib import Path
from typing import Any

__all__ = [
    "COLUMN_MAP",
    "SBICode",
    "load_dictionary",
    "normalize_code",
    "read_dictionary_csv",
    "read_dictionary_excel",
    "write_dictionary_csv",
]

#: Mapping of the original (Dutch) workbook headers to :class:`SBICode` fields.
COLUMN_MAP: dict[str, str] = {
    "code": "code",
    "description": "description",
    "Warning": "warning",
    "Tijdseries in deze SBI": "n_timeseries",
    "Profiel hoogste correlatie": "best_profile",
    "Standaard profiel": "standard_profile",
    "correlatie geaggregeerde profiel": "corr_best_profile",
    "correlatie standaard profiel": "corr_standard_profile",
    "KO_profiel": "ko_profile",
    "Profiel aanwezig": "has_profile",
}

_MISSING = {"", "na", "nan", "none", "null", "0"}
_BUNDLED_CSV = "sbi_dictionary.csv"


@dataclass(frozen=True)
class SBICode:
    """One row of the SBI dictionary.

    Attributes:
        code: SBI code as a string, e.g. ``"8411"`` (leading zeros preserved).
        description: Dutch description of the economic activity.
        warning: Reliability warning (set when fewer than 10 time series
            were available for this code), otherwise ``None``.
        n_timeseries: Number of metered connections used to build the profile.
        best_profile: Profile with the highest correlation. Either an SBI code
            (whose profile is in the database) or a standard ``KO_*`` profile.
        standard_profile: The standard ``KO_*`` profile for this sector.
        corr_best_profile: Correlation with the aggregated ``best_profile``.
        corr_standard_profile: Correlation with the ``standard_profile``.
        ko_profile: Recommended standard (``KO_*``) profile.
        has_profile: ``True`` if a measured profile exists for this exact code.
    """

    code: str
    description: str | None = None
    warning: str | None = None
    n_timeseries: int | None = None
    best_profile: str | None = None
    standard_profile: str | None = None
    corr_best_profile: float | None = None
    corr_standard_profile: float | None = None
    ko_profile: str | None = None
    has_profile: bool = False

    @property
    def level(self) -> int:
        """Depth in the SBI hierarchy (number of digits in the code)."""
        return len(self.code)

    @property
    def parent_code(self) -> str | None:
        """The code one level up in the hierarchy (``"8411"`` -> ``"841"``)."""
        return self.code[:-1] if len(self.code) > 2 else None

    @property
    def recommended_profile(self) -> str | None:
        """The recommended profile: ``best_profile``, falling back to ``ko_profile``."""
        return self.best_profile or self.ko_profile

    @property
    def recommends_measured_profile(self) -> bool:
        """``True`` if the recommendation is a measured SBI profile (not ``KO_*``)."""
        rec = self.recommended_profile
        return rec is not None and rec.isdigit()

    def to_dict(self) -> dict[str, Any]:
        """Return the record as a plain dictionary."""
        return asdict(self)


def normalize_code(value: Any) -> str:
    """Normalise an SBI code to its canonical string form.

    SBI divisions always have at least two digits, but spreadsheets tend to
    drop leading zeros (``2`` instead of ``02``).  Integers, floats such as
    ``1.0`` and padded strings are all handled.

    Examples:
        >>> normalize_code("0119")
        '0119'
        >>> normalize_code(2)
        '02'
        >>> normalize_code(" 8411 ")
        '8411'
    """
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    code = str(value).strip()
    if code.endswith(".csv"):
        code = code[:-4]
    if not code.isdigit():
        raise ValueError(f"Invalid SBI code: {value!r}")
    return code.zfill(2)


def _clean_str(value: Any) -> str | None:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    text = str(value).strip()
    return None if text.lower() in _MISSING else text


def _clean_float(value: Any) -> float | None:
    if value is None:
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return None if math.isnan(number) else number


def _clean_int(value: Any) -> int | None:
    number = _clean_float(value)
    return None if number is None else int(number)


def _clean_profile(value: Any) -> str | None:
    text = _clean_str(value)
    if text is not None and text.isdigit():
        return normalize_code(text)
    return text


def _clean_corr(value: Any) -> float | None:
    # The workbook uses "NA" and an exact 0 as placeholders for "no correlation".
    number = _clean_float(value)
    return None if number == 0 else number


def _record_from_mapping(row: dict[str, Any]) -> SBICode:
    return SBICode(
        code=normalize_code(row["code"]),
        description=_clean_str(row.get("description")),
        warning=_clean_str(row.get("warning")),
        n_timeseries=_clean_int(row.get("n_timeseries")),
        best_profile=_clean_profile(row.get("best_profile")),
        standard_profile=_clean_profile(row.get("standard_profile")),
        corr_best_profile=_clean_corr(row.get("corr_best_profile")),
        corr_standard_profile=_clean_corr(row.get("corr_standard_profile")),
        ko_profile=_clean_profile(row.get("ko_profile")),
        has_profile=bool(_clean_int(row.get("has_profile")) or 0),
    )


def _check_unique(records: list[SBICode]) -> list[SBICode]:
    seen: set[str] = set()
    for rec in records:
        if rec.code in seen:
            raise ValueError(f"Duplicate SBI code in dictionary: {rec.code}")
        seen.add(rec.code)
    return records


def read_dictionary_excel(path: str | Path, sheet: str | int = 0) -> list[SBICode]:
    """Parse the original Liander Excel workbook.

    Requires the optional ``openpyxl`` dependency
    (``pip install "liander-open-data[excel]"``).

    Args:
        path: Path to ``Overzicht met per SBI-code een aanbevolen profiel.xlsx``.
        sheet: Sheet name or zero-based index (default: the first sheet).

    Returns:
        One :class:`SBICode` per row.
    """
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise ImportError(
            'Reading the Excel dictionary requires openpyxl: pip install "liander-open-data[excel]"'
        ) from exc

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if isinstance(sheet, str) else wb.worksheets[sheet]
        rows = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows)]
        missing = [h for h in COLUMN_MAP if h not in header]
        if missing:
            raise ValueError(f"Dictionary sheet is missing columns: {missing}")
        index = {COLUMN_MAP[h]: header.index(h) for h in COLUMN_MAP}
        records = [
            _record_from_mapping({key: row[i] for key, i in index.items()})
            for row in rows
            if row and row[index["code"]] is not None
        ]
    finally:
        wb.close()
    return _check_unique(records)


def read_dictionary_csv(path: str | Path) -> list[SBICode]:
    """Read a dictionary previously written by :func:`write_dictionary_csv`."""
    with open(path, newline="", encoding="utf-8") as fh:
        return _check_unique([_record_from_mapping(row) for row in csv.DictReader(fh)])


def write_dictionary_csv(records: Iterable[SBICode], path: str | Path) -> None:
    """Write dictionary records to a UTF-8 CSV file with English headers."""
    names = [f.name for f in fields(SBICode)]
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=names)
        writer.writeheader()
        for rec in records:
            row = rec.to_dict()
            row["has_profile"] = int(rec.has_profile)
            writer.writerow(row)


def load_dictionary(path: str | Path | None = None) -> list[SBICode]:
    """Load the SBI dictionary.

    Args:
        path: Optional path to an ``.xlsx`` workbook or a ``.csv`` export.
            When omitted, the copy bundled with the package is used.

    Returns:
        A list of :class:`SBICode` records.
    """
    if path is None:
        source = resources.files("liander_open_data").joinpath("data", _BUNDLED_CSV)
        with resources.as_file(source) as bundled:
            return read_dictionary_csv(bundled)
    path = Path(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return read_dictionary_excel(path)
    return read_dictionary_csv(path)
